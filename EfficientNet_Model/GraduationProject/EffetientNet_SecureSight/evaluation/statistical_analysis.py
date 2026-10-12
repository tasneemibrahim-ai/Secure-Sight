import os
import json
import logging
import datetime
from typing import Dict, List, Any

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from scipy import stats
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

METRIC_COLUMNS = ['accuracy', 'precision', 'recall', 'f1_score', 'auc', 'loss', 'latency']


def _safe_mkdir(path: str) -> None:
    os.makedirs(path, exist_ok=True)


def setup_logging(log_path: str) -> None:
    logging.basicConfig(
        filename=log_path,
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
    )


def create_directories(base_path: str) -> Dict[str, str]:
    subdirs = {
        'raw_data': os.path.join(base_path, 'raw_data'),
        'summary': os.path.join(base_path, 'summary'),
        'boxplots': os.path.join(base_path, 'plots', 'boxplots'),
        'lineplots': os.path.join(base_path, 'plots', 'lineplots'),
        'barplots': os.path.join(base_path, 'plots', 'barplots'),
        'heatmaps': os.path.join(base_path, 'plots', 'heatmaps'),
        'histograms': os.path.join(base_path, 'plots', 'histograms'),
        'reports': os.path.join(base_path, 'reports'),
        'logs': os.path.join(base_path, 'logs'),
    }
    for path in subdirs.values():
        _safe_mkdir(path)
    return subdirs


def load_evaluation_results(results_path: str) -> pd.DataFrame:
    detailed_dir = os.path.join(results_path, 'detailed_batches')
    if not os.path.isdir(detailed_dir):
        raise FileNotFoundError(f'Could not find detailed_batches in {results_path}')

    rows: List[Dict[str, Any]] = []
    for batch_dir in sorted(os.listdir(detailed_dir)):
        batch_path = os.path.join(detailed_dir, batch_dir)
        if not os.path.isdir(batch_path):
            continue
        batch_id = batch_dir.replace('batch_', '')
        for run_dir in sorted(os.listdir(batch_path)):
            run_path = os.path.join(batch_path, run_dir)
            summary_path = os.path.join(run_path, 'run_summary.json')
            if not os.path.exists(summary_path):
                continue
            with open(summary_path, 'r', encoding='utf-8') as f:
                summary = json.load(f)
            final_metrics = summary.get('final_metrics', {}) or {}
            loss = final_metrics.get('loss') if final_metrics.get('loss') is not None else summary.get('loss')
            rows.append({
                'batch_id': batch_id,
                'run_id': run_dir,
                'accuracy': final_metrics.get('accuracy', np.nan),
                'precision': final_metrics.get('precision', np.nan),
                'recall': final_metrics.get('recall', np.nan),
                'f1_score': final_metrics.get('f1', np.nan),
                'auc': final_metrics.get('auc', np.nan),
                'loss': loss if loss is not None else np.nan,
                'latency': summary.get('latency', {}).get('mean_ms', np.nan),
            })

    df = pd.DataFrame(rows)
    for col in ['loss']:
        if col not in df.columns:
            df[col] = np.nan
    return df


def _summarize_series(series: pd.Series) -> Dict[str, float]:
    values = series.dropna().astype(float)
    if values.empty:
        return {k: float('nan') for k in ['mean', 'std', 'variance', 'min', 'max', 'median', 'count']}
    return {
        'mean': float(values.mean()),
        'std': float(values.std()),
        'variance': float(values.var()),
        'min': float(values.min()),
        'max': float(values.max()),
        'median': float(values.median()),
        'count': int(values.count()),
    }


def compute_global_statistics(df: pd.DataFrame) -> Dict[str, Dict[str, float]]:
    return {metric: _summarize_series(df[metric]) for metric in METRIC_COLUMNS if metric in df.columns}


def compute_per_batch_statistics(df: pd.DataFrame) -> Dict[str, Dict[str, Dict[str, float]]]:
    per_batch: Dict[str, Dict[str, Dict[str, float]]] = {}
    for batch_id, batch_df in df.groupby('batch_id'):
        per_batch[batch_id] = {
            metric: _summarize_series(batch_df[metric])
            for metric in METRIC_COLUMNS if metric in batch_df.columns
        }
    return per_batch


def compute_confidence_intervals(df: pd.DataFrame, confidence: float = 0.95) -> Dict[str, Dict[str, float]]:
    cis: Dict[str, Dict[str, float]] = {}
    for metric in METRIC_COLUMNS:
        if metric not in df.columns:
            continue
        series = df[metric].dropna().astype(float)
        if series.empty:
            cis[metric] = {'mean': float('nan'), 'ci_lower': float('nan'), 'ci_upper': float('nan')}
            continue
        mean = series.mean()
        if len(series) < 2:
            ci_lower = ci_upper = mean
        else:
            sem = stats.sem(series)
            ci_lower, ci_upper = stats.t.interval(confidence, len(series) - 1, loc=mean, scale=sem)
        cis[metric] = {
            'mean': float(mean),
            'ci_lower': float(ci_lower),
            'ci_upper': float(ci_upper),
        }
    return cis


def compute_correlation_matrix(df: pd.DataFrame) -> Dict[str, Dict[str, float]]:
    metrics = [metric for metric in METRIC_COLUMNS if metric in df.columns]
    corr = df[metrics].corr()
    return corr.fillna(0).to_dict()


def compute_stability_analysis(df: pd.DataFrame) -> Dict[str, Any]:
    stability: Dict[str, Any] = {}
    if 'batch_id' not in df.columns:
        return stability
    for batch_id, batch_df in df.groupby('batch_id'):
        batch_stats = {
            metric: {
                'std': float(batch_df[metric].dropna().std()) if metric in batch_df else float('nan'),
                'variance': float(batch_df[metric].dropna().var()) if metric in batch_df else float('nan'),
            }
            for metric in METRIC_COLUMNS if metric in batch_df.columns
        }
        instability_score = np.nanmean([batch_stats[m]['std'] for m in batch_stats if not np.isnan(batch_stats[m]['std'])])
        stability[batch_id] = {
            'metrics': batch_stats,
            'instability_score': float(instability_score) if not np.isnan(instability_score) else float('nan'),
        }
    return stability


def find_best_worst_runs(df: pd.DataFrame) -> Dict[str, Dict[str, Any]]:
    if df.empty:
        return {'best': {}, 'worst': {}}
    best_row = df.loc[df['f1_score'].idxmax()].to_dict()
    worst_row = df.loc[df['f1_score'].idxmin()].to_dict()
    return {
        'best': {k: (float(v) if isinstance(v, (int, float, np.generic)) else v) for k, v in best_row.items()},
        'worst': {k: (float(v) if isinstance(v, (int, float, np.generic)) else v) for k, v in worst_row.items()},
    }


def _save_figure(fig: plt.Figure, path: str) -> None:
    _safe_mkdir(os.path.dirname(path))
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)


def plot_boxplots(df: pd.DataFrame, target_dir: str) -> None:
    metrics = [m for m in METRIC_COLUMNS if m in df.columns and m != 'latency']
    for metric in metrics:
        fig, ax = plt.subplots(figsize=(10, 6))
        sns.boxplot(x='batch_id', y=metric, data=df, palette='Set2', ax=ax)
        ax.set_title(f'{metric.capitalize()} Distribution by Batch', fontsize=14)
        ax.set_xlabel('Batch ID')
        ax.set_ylabel(metric.capitalize())
        _save_figure(fig, os.path.join(target_dir, 'boxplots', f'{metric}_boxplot.png'))


def plot_histograms(df: pd.DataFrame, target_dir: str) -> None:
    metrics = [m for m in METRIC_COLUMNS if m in df.columns]
    for metric in metrics:
        fig, ax = plt.subplots(figsize=(10, 6))
        sns.histplot(df[metric].dropna(), kde=True, color='#4c78a8', ax=ax)
        ax.set_title(f'{metric.capitalize()} Distribution', fontsize=14)
        ax.set_xlabel(metric.capitalize())
        ax.set_ylabel('Frequency')
        _save_figure(fig, os.path.join(target_dir, 'histograms', f'{metric}_histogram.png'))


def plot_lineplots(df: pd.DataFrame, target_dir: str) -> None:
    metrics = [m for m in METRIC_COLUMNS if m in df.columns and m != 'loss']
    df_sorted = df.sort_values(['batch_id', 'run_id'])
    run_index = list(range(1, len(df_sorted) + 1))
    for metric in metrics:
        fig, ax = plt.subplots(figsize=(10, 6))
        ax.plot(run_index, df_sorted[metric].astype(float).values, marker='o', linewidth=2)
        ax.set_title(f'{metric.capitalize()} Trend Across Runs', fontsize=14)
        ax.set_xlabel('Run Index')
        ax.set_ylabel(metric.capitalize())
        ax.grid(alpha=0.3)
        _save_figure(fig, os.path.join(target_dir, 'lineplots', f'{metric}_lineplot.png'))
    if 'accuracy' in df.columns and 'f1_score' in df.columns:
        fig, ax = plt.subplots(figsize=(10, 6))
        sns.scatterplot(x='accuracy', y='f1_score', hue='batch_id', data=df, palette='Set1', ax=ax)
        ax.set_title('Accuracy vs F1 Score by Batch', fontsize=14)
        ax.set_xlabel('Accuracy')
        ax.set_ylabel('F1 Score')
        _save_figure(fig, os.path.join(target_dir, 'lineplots', 'accuracy_vs_f1_scatter.png'))


def plot_barplots(df: pd.DataFrame, target_dir: str) -> None:
    metrics = [m for m in METRIC_COLUMNS if m in df.columns]
    batch_stats = df.groupby('batch_id')[metrics].agg(['mean', 'std']).dropna(axis=1, how='all')
    for metric in metrics:
        if metric not in batch_stats.columns:
            continue
        means = batch_stats[(metric, 'mean')]
        errors = batch_stats[(metric, 'std')]
        fig, ax = plt.subplots(figsize=(10, 6))
        ax.bar(means.index.astype(str), means.values, yerr=errors.values, capsize=6, color='#2a9d8f')
        ax.set_title(f'{metric.capitalize()} Mean ± Std by Batch', fontsize=14)
        ax.set_xlabel('Batch ID')
        ax.set_ylabel(metric.capitalize())
        ax.grid(alpha=0.3, axis='y')
        _save_figure(fig, os.path.join(target_dir, 'barplots', f'{metric}_barplot.png'))


def plot_heatmaps(df: pd.DataFrame, target_dir: str) -> None:
    metrics = [m for m in METRIC_COLUMNS if m in df.columns]
    corr = df[metrics].corr()
    fig, ax = plt.subplots(figsize=(10, 8))
    sns.heatmap(corr, annot=True, cmap='coolwarm', fmt='.2f', ax=ax)
    ax.set_title('Correlation Matrix', fontsize=14)
    _save_figure(fig, os.path.join(target_dir, 'heatmaps', 'correlation_heatmap.png'))


def _write_table(document: Document, headers: List[str], rows: List[List[Any]], style: str = 'Light List Accent 1') -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = style
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = str(header)
    for row in rows:
        row_cells = table.add_row().cells
        for idx, cell in enumerate(row):
            row_cells[idx].text = str(cell)
    document.add_paragraph('')


def _add_image(document: Document, image_path: str, caption: str) -> None:
    if not os.path.exists(image_path):
        return
    document.add_paragraph(caption, style='Intense Quote')
    document.add_picture(image_path, width=Inches(6.5))
    last_paragraph = document.paragraphs[-1]
    last_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    document.add_paragraph('')


def generate_word_report(
    df: pd.DataFrame,
    global_stats: Dict[str, Dict[str, float]],
    cis: Dict[str, Dict[str, float]],
    per_batch: Dict[str, Dict[str, Dict[str, float]]],
    correlation_matrix: Dict[str, Dict[str, float]],
    stability: Dict[str, Any],
    best_worst: Dict[str, Dict[str, Any]],
    paths: Dict[str, str],
) -> str:
    doc = Document()
    doc.core_properties.title = 'SecureSight Statistical Analysis Report'
    doc.core_properties.author = 'SecureSight Evaluation System'
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run('SecureSight\n')
    run.bold = True
    run.font.size = Pt(28)
    run = title.add_run('Statistical Analysis Report\n')
    run.font.size = Pt(20)
    run = title.add_run(datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    run.font.size = Pt(12)
    doc.add_page_break()

    doc.add_heading('Overview', level=1)
    doc.add_paragraph('This report summarizes the SecureSight evaluation results, combining batch and run metrics into a centralized statistical overview.')
    doc.add_paragraph(f'Total Runs: {len(df)}')
    doc.add_paragraph(f'Total Batches: {df["batch_id"].nunique() if "batch_id" in df.columns else 0}')
    doc.add_paragraph('Evaluated metrics include accuracy, precision, recall, F1 score, AUC, loss, and latency where available.')

    doc.add_heading('Global Statistical Analysis', level=1)
    headers = ['Metric', 'Mean', 'Std', 'Variance', 'Min', 'Max', 'Median', 'Count']
    rows = [
        [
            metric,
            f'{stats["mean"]:.4f}',
            f'{stats["std"]:.4f}',
            f'{stats["variance"]:.4f}',
            f'{stats["min"]:.4f}',
            f'{stats["max"]:.4f}',
            f'{stats["median"]:.4f}',
            str(stats.get('count', ''))
        ]
        for metric, stats in global_stats.items()
    ]
    _write_table(doc, headers, rows)

    doc.add_heading('Confidence Intervals', level=2)
    headers = ['Metric', 'Mean', 'CI Lower', 'CI Upper']
    rows = [
        [metric, f'{ci["mean"]:.4f}', f'{ci["ci_lower"]:.4f}', f'{ci["ci_upper"]:.4f}']
        for metric, ci in cis.items()
    ]
    _write_table(doc, headers, rows)

    doc.add_heading('Per-Batch Analysis', level=1)
    for batch_id, metrics in per_batch.items():
        doc.add_heading(f'Batch {batch_id}', level=2)
        rows = [
            [
                metric,
                f'{values["mean"]:.4f}',
                f'{values["std"]:.4f}',
                f'{values["variance"]:.4f}',
                f'{values["min"]:.4f}',
                f'{values["max"]:.4f}',
                f'{values.get("median", float("nan")):.4f}' if not pd.isna(values.get('median', float('nan'))) else 'nan',
                str(values.get('count', ''))
            ]
            for metric, values in metrics.items()
        ]
        _write_table(doc, ['Metric', 'Mean', 'Std', 'Variance', 'Min', 'Max', 'Median', 'Count'], rows)
        stability_score = stability.get(batch_id, {}).get('instability_score', float('nan'))
        doc.add_paragraph(f'Batch stability score: {stability_score:.4f}')

    doc.add_heading('Best and Worst Runs', level=1)
    if best_worst.get('best'):
        doc.add_paragraph('Best Run')
        rows = [[k, best_worst['best'][k]] for k in best_worst['best'] if k != 'batch_id' and k != 'run_id']
        _write_table(doc, ['Metric', 'Value'], [['batch_id', best_worst['best']['batch_id']], ['run_id', best_worst['best']['run_id']]] + rows)
    if best_worst.get('worst'):
        doc.add_paragraph('Worst Run')
        rows = [[k, best_worst['worst'][k]] for k in best_worst['worst'] if k != 'batch_id' and k != 'run_id']
        _write_table(doc, ['Metric', 'Value'], [['batch_id', best_worst['worst']['batch_id']], ['run_id', best_worst['worst']['run_id']]] + rows)

    doc.add_heading('Correlation Analysis', level=1)
    corr_headers = [''] + list(correlation_matrix.keys())
    corr_rows = []
    for metric, row in correlation_matrix.items():
        corr_rows.append([metric] + [f'{row[col]:.4f}' for col in correlation_matrix.keys()])
    _write_table(doc, corr_headers, corr_rows)
    _add_image(doc, os.path.join(paths['heatmaps'], 'correlation_heatmap.png'), 'Figure: Correlation Matrix Heatmap')

    doc.add_heading('Visualizations', level=1)
    for folder_key in ['boxplots', 'lineplots', 'barplots', 'heatmaps', 'histograms']:
        subdir = paths[folder_key]
        doc.add_heading(folder_key.replace('plots', '').capitalize(), level=2)
        for image_name in sorted(os.listdir(subdir)):
            image_path = os.path.join(subdir, image_name)
            if image_path.lower().endswith('.png'):
                caption = f'Figure: {image_name.replace("_", " ").replace(".png", "").title()}'
                _add_image(doc, image_path, caption)

    doc.add_heading('Final Insights', level=1)
    overall_best_metric = max(global_stats.items(), key=lambda item: item[1]['mean'])[0]
    overall_worst_metric = min(global_stats.items(), key=lambda item: item[1]['mean'])[0]
    doc.add_paragraph(f'The strongest overall metric is {overall_best_metric} with a mean value of {global_stats[overall_best_metric]["mean"]:.4f}.')
    doc.add_paragraph(f'The weakest overall metric is {overall_worst_metric} with a mean value of {global_stats[overall_worst_metric]["mean"]:.4f}.')
    high_instability = sorted(stability.items(), key=lambda item: item[1].get('instability_score', 0), reverse=True)
    if high_instability:
        doc.add_paragraph(f'The most unstable batch is Batch {high_instability[0][0]} with instability score {high_instability[0][1]["instability_score"]:.4f}.')
    doc.add_paragraph('The report highlights batch-level consistency, run-level variability, and the correlation of metrics to support deployment decisions.')

    output_docx = os.path.join(paths['reports'], 'statistical_analysis_report.docx')
    doc.save(output_docx)
    return output_docx


def generate_text_report(
    paths: Dict[str, str],
    global_stats: Dict[str, Dict[str, float]],
    cis: Dict[str, Dict[str, float]],
    stability: Dict[str, Any],
    best_worst: Dict[str, Dict[str, Any]],
) -> None:
    text_path = os.path.join(paths['reports'], 'report.txt')
    lines: List[str] = [
        'SecureSight Statistical Analysis Report',
        '=' * 60,
        '\nGlobal Metrics:',
    ]
    for metric, stats in global_stats.items():
        lines.append(f'  {metric}: {stats["mean"]:.4f} ± {stats["std"]:.4f} (min: {stats["min"]:.4f}, max: {stats["max"]:.4f}, median: {stats["median"]:.4f})')
    lines.append('\nConfidence Intervals (95%):')
    for metric, ci in cis.items():
        lines.append(f'  {metric}: [{ci["ci_lower"]:.4f}, {ci["ci_upper"]:.4f}]')
    lines.append('\nStability Analysis:')
    for batch_id, stats in stability.items():
        lines.append(f'  Batch {batch_id}: instability_score={stats["instability_score"]:.4f}')
    lines.append('\nBest Run:')
    lines.append(f'  Batch {best_worst["best"]["batch_id"]}, Run {best_worst["best"]["run_id"]}, F1={best_worst["best"]["f1_score"]:.4f}')
    lines.append('\nWorst Run:')
    lines.append(f'  Batch {best_worst["worst"]["batch_id"]}, Run {best_worst["worst"]["run_id"]}, F1={best_worst["worst"]["f1_score"]:.4f}')
    with open(text_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))


def generate_json_report(paths: Dict[str, str], global_stats: Dict[str, Dict[str, float]], cis: Dict[str, Dict[str, float]], per_batch: Dict[str, Dict[str, Dict[str, float]]], stability: Dict[str, Any], best_worst: Dict[str, Dict[str, Any]]) -> None:
    output = {
        'global_metrics': global_stats,
        'confidence_intervals': cis,
        'per_batch_metrics': per_batch,
        'stability_analysis': stability,
        'best_worst_runs': best_worst,
    }
    with open(os.path.join(paths['reports'], 'report.json'), 'w', encoding='utf-8') as f:
        json.dump(output, f, indent=2)


def run_full_statistical_analysis(results_path: str) -> None:
    base_path = os.path.join(results_path, 'statistical_analysis')
    paths = create_directories(base_path)
    log_path = os.path.join(paths['logs'], 'analysis_log.txt')
    setup_logging(log_path)
    logging.info('Starting statistical analysis')

    df = load_evaluation_results(results_path)
    if df.empty:
        raise ValueError('No evaluation runs were loaded from the results path.')

    raw_csv_path = os.path.join(paths['raw_data'], 'all_runs.csv')
    df.to_csv(raw_csv_path, index=False)
    logging.info(f'Raw data saved to {raw_csv_path}')

    global_stats = compute_global_statistics(df)
    per_batch = compute_per_batch_statistics(df)
    cis = compute_confidence_intervals(df)
    correlation_matrix = compute_correlation_matrix(df)
    stability = compute_stability_analysis(df)
    best_worst = find_best_worst_runs(df)

    with open(os.path.join(paths['summary'], 'global_metrics.json'), 'w', encoding='utf-8') as f:
        json.dump(global_stats, f, indent=2)
    with open(os.path.join(paths['summary'], 'per_batch_metrics.json'), 'w', encoding='utf-8') as f:
        json.dump(per_batch, f, indent=2)
    with open(os.path.join(paths['summary'], 'confidence_intervals.json'), 'w', encoding='utf-8') as f:
        json.dump(cis, f, indent=2)
    with open(os.path.join(paths['summary'], 'correlation_matrix.json'), 'w', encoding='utf-8') as f:
        json.dump(correlation_matrix, f, indent=2)
    with open(os.path.join(paths['summary'], 'best_worst_runs.json'), 'w', encoding='utf-8') as f:
        json.dump(best_worst, f, indent=2)
    logging.info('Summary JSON files created')

    plot_boxplots(df, paths['boxplots'])
    plot_histograms(df, paths['histograms'])
    plot_lineplots(df, paths['lineplots'])
    plot_barplots(df, paths['barplots'])
    plot_heatmaps(df, paths['heatmaps'])
    logging.info('Plots generated')

    word_path = generate_word_report(df, global_stats, cis, per_batch, correlation_matrix, stability, best_worst, paths)
    logging.info(f'Word report created at {word_path}')

    generate_text_report(paths, global_stats, cis, stability, best_worst)
    generate_json_report(paths, global_stats, cis, per_batch, stability, best_worst)
    logging.info('Text and JSON reports created')
    logging.info('Statistical analysis complete')


if __name__ == '__main__':
    run_full_statistical_analysis(os.path.join(os.path.dirname(__file__), 'evaluation'))